#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""step4_word_service —— Step4 的图文 Word 文档与批注（officeCLI）。

设计（用户架构）：
  * Step1-3 做数据清洗/基于语料生成论文；
  * **Step4 背后调用 officeCLI**：把初稿落成**图文 Word**（标题/各节正文/图/参考文献）；
  * **Step4 的每轮交互 = 一条 Word 批注（comment）**，记录进 .docx；
  * Step5 导出 Word（复制交付）。

officeCLI 定位：`MIMW_OFFICECLI`（officecli.js 路径）→ npm-global 默认；
node 取 PATH。所有子进程强制 UTF-8 解码 + 无窗口（GUI 下不弹控制台）。
每次操作后 `save` + `close`：保证外部程序（Word/WPS）能读到最新文件。
"""

from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import threading
import time
from pathlib import Path
from typing import Any, Dict, List, Optional

from src.utils.logger import get_logger

logger = get_logger(__name__)

_NO_WINDOW = 0x08000000 if os.name == "nt" else 0
_SECTION_ORDER = (
    ("abstract", "Abstract"),
    ("intro", "Introduction"),
    ("methods", "Methods"),
    ("results", "Results"),
    ("discussion", "Discussion"),
    ("conclusion", "Conclusion"),
)


# ---------------------------------------------------------------- 可执行定位

def _officecli_js() -> str:
    env = os.environ.get("MIMW_OFFICECLI", "").strip()
    cands: List[Path] = []
    if env:
        cands.append(Path(env))
    for base in (os.environ.get("LOCALAPPDATA"), os.environ.get("APPDATA")):
        if base:
            cands.append(Path(base) / "npm-global" / "node_modules" /
                         "@officecli" / "officecli" / "officecli.js")
            cands.append(Path(base) / "npm" / "node_modules" /
                         "@officecli" / "officecli" / "officecli.js")
    for c in cands:
        try:
            if c.is_file():
                return str(c)
        except Exception:
            continue
    return ""


def available() -> bool:
    return bool(shutil.which("node")) and bool(_officecli_js())


def _run(*args: str, timeout: float = 240.0) -> Dict[str, Any]:
    node = shutil.which("node")
    js = _officecli_js()
    if not node or not js:
        return {"ok": False, "stdout": "", "stderr": "officecli not available"}
    try:
        r = subprocess.run(
            [node, js, *[str(a) for a in args]],
            capture_output=True, text=True,
            encoding="utf-8", errors="replace",
            timeout=timeout, creationflags=_NO_WINDOW,
        )
        return {"ok": r.returncode == 0, "rc": r.returncode,
                "stdout": r.stdout or "", "stderr": r.stderr or ""}
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "stdout": "", "stderr": f"{type(e).__name__}: {e}"}


def _ok(res: Dict[str, Any], what: str) -> bool:
    if res.get("ok"):
        return True
    logger.warning("officecli %s failed: %s | %s", what,
                   str(res.get("stderr") or "")[:300], str(res.get("stdout") or "")[:200])
    return False


# ---------------------------------------------------------------- 存储路径

def _storage_dir(project_id: str) -> str:
    try:
        from src.services.database_service import DatabaseService
        from src.services.project_storage_service import ProjectStorageService
        svc = ProjectStorageService(DatabaseService())
        return str(svc.get_project_storage_path(str(project_id)))
    except Exception:
        return ""


def step4_dir(project_id: str, out_dir: str = "") -> str:
    if out_dir:
        d = str(out_dir)
    else:
        base = _storage_dir(project_id)
        if not base:
            return ""
        d = os.path.join(base, "exports", "step4")
    try:
        os.makedirs(d, exist_ok=True)
    except Exception:
        return ""
    return d


def _sidecar_path(d: str) -> str:
    return os.path.join(d, "latest_step4_word.json")


def _slug(title: str) -> str:
    s = re.sub(r"[\\/:*?\"<>|\s]+", "_", str(title or "draft")).strip("_")
    return (s or "draft")[:40]


# ---------------------------------------------------------------- 构建 docx

def _split_blocks(text: str) -> List[str]:
    t = str(text or "").replace("\r\n", "\n")
    blocks = [b.strip() for b in re.split(r"\n{2,}", t) if b.strip()]
    return blocks


def _add_paragraph(docx: str, text: str, *, style: str = "", bold: bool = False,
                   size: str = "", align: str = "") -> bool:
    if not text:
        return False
    # \n → 软换行（\v），保持单段落（officecli: \n 会拆段落）
    one = str(text).replace("\n", "\x0b")
    args = ["add", docx, "/body", "--type", "paragraph", "--prop", f"text={one}"]
    if style:
        args += ["--prop", f"style={style}"]
    if bold:
        args += ["--prop", "bold=true"]
    if size:
        args += ["--prop", f"size={size}"]
    if align:
        args += ["--prop", f"align={align}"]
    return _ok(_run(*args), "add paragraph")


def _add_picture(docx: str, image_path: str, caption: str = "") -> bool:
    r = _run("add", docx, "/body", "--type", "paragraph")
    if not _ok(r, "add empty paragraph"):
        return False
    m = re.search(r"p\[(\d+)\]|@paraId=([0-9A-Fa-f]+)", r.get("stdout") or "")
    para = "/body/p[1]"  # 兜底
    if m:
        para = (f"/body/p[{m.group(1)}]" if m.group(1)
                else f"/body/p[@paraId={m.group(2)}]")
    ok = _ok(_run("add", docx, para, "--type", "picture",
                  "--prop", f"src={str(image_path)}",
                  "--prop", "width=14cm",
                  "--prop", f"alt={caption or Path(image_path).name}"), "add picture")
    if ok and caption:
        _add_paragraph(docx, caption, size="9", align="center")
    return ok


def build_step4_docx(project_id: str, result_dict: Dict[str, Any],
                     out_dir: str = "", image_paths: Optional[List[str]] = None) -> Optional[str]:
    """由 Step3/4 结果生成**图文 Word**；返回 docx 路径（失败 None）。"""
    if not available():
        logger.warning("officecli 不可用（node/officecli.js 未找到），跳过 Step4 Word 生成")
        return None
    d = step4_dir(project_id, out_dir=out_dir)
    if not d:
        return None

    result = result_dict or {}
    outline = result.get("outline") or {}
    units = result.get("units") or {}
    refs = result.get("references") or []
    title = str(outline.get("title") or "Draft").strip() or "Draft"

    docx = os.path.join(d, f"{_slug(title)}__{time.strftime('%Y%m%d_%H%M%S')}.docx")
    if not _ok(_run("create", docx, timeout=120), "create"):
        return None

    _add_paragraph(docx, title, style="Heading1", bold=True, size="20", align="center")

    for key, name in _SECTION_ORDER:
        content = ""
        u = units.get(key)
        if isinstance(u, dict):
            content = str(u.get("content") or "")
        blocks = _split_blocks(content)
        if not blocks:
            continue
        _add_paragraph(docx, name, style="Heading1", bold=True, size="14")
        for b in blocks:
            _add_paragraph(docx, b, size="11")
        if key == "results":
            _insert_figures(docx, project_id, image_paths)

    # 若 Results 不存在，图片仍附在文末
    if "results" not in {k for k, _ in _SECTION_ORDER if (units.get(k) or {}).get("content")}:
        _insert_figures(docx, project_id, image_paths)

    if refs:
        _add_paragraph(docx, "References", style="Heading1", bold=True, size="14")
        for i, ref in enumerate(refs[:80], 1):
            cit = ""
            if isinstance(ref, dict):
                cit = str(ref.get("citation") or "")
            elif isinstance(ref, str):
                cit = ref
            if cit.strip():
                _add_paragraph(docx, f"[{i}] {cit.strip()}", size="10")

    _ok(_run("validate", docx, timeout=120), "validate")
    _ok(_run("save", docx, timeout=120), "save")
    _ok(_run("close", docx, timeout=120), "close")

    try:
        with open(_sidecar_path(d), "w", encoding="utf-8") as f:
            json.dump({"path": docx, "title": title, "project_id": str(project_id),
                       "ts": time.strftime("%Y-%m-%d %H:%M:%S")}, f,
                      ensure_ascii=False, indent=2)
    except Exception:
        pass
    return docx


def _insert_figures(docx: str, project_id: str, image_paths: Optional[List[str]]) -> None:
    imgs: List[str] = []
    if image_paths:
        imgs = [str(p) for p in image_paths if str(p) and os.path.isfile(str(p))]
    else:
        base = _storage_dir(project_id)
        if base:
            cand: List[str] = []
            for root, _dirs, files in os.walk(base):
                for fn in files:
                    if fn.lower().endswith((".png", ".jpg", ".jpeg")):
                        cand.append(os.path.join(root, fn))
            cand.sort(key=lambda p: os.path.getmtime(p), reverse=True)
            imgs = cand[:6]
    for i, img in enumerate(imgs[:8], 1):
        _add_picture(docx, img, caption=f"图 {i}. {Path(img).name}")


# ---------------------------------------------------------------- 批注

def _find_anchor(docx: str, anchor_text: str) -> str:
    needle = re.sub(r"\s+", "", str(anchor_text or ""))[:40]
    if not needle:
        return "/body/p[1]"
    r = _run("query", docx, "paragraph", "--json", timeout=120)
    if not r.get("ok"):
        return "/body/p[1]"
    try:
        data = json.loads(r.get("stdout") or "{}")
        results = ((data.get("data") or {}).get("results")) or []
        for item in results:
            txt = re.sub(r"\s+", "", str(item.get("text") or ""))
            if needle and needle in txt:
                return str(item.get("path") or "/body/p[1]")
    except Exception:
        pass
    return "/body/p[1]"


def add_feedback_comment(docx_path: str, text: str, *, author: str = "研究者",
                         anchor_text: str = "", parent_id: Optional[int] = None) -> bool:
    """把一轮交互/反馈追加为 Word 批注（记录进 .docx）。"""
    if not available() or not docx_path or not os.path.isfile(docx_path):
        return False
    body = str(text or "").strip()
    if not body:
        return False
    anchor = _find_anchor(docx_path, anchor_text)
    args = ["add", docx_path, anchor, "--type", "comment",
            "--prop", f"author={author}", "--prop", f"text={body[:1800]}"]
    if parent_id is not None:
        args += ["--prop", f"parentId={int(parent_id)}"]
    ok = _ok(_run(*args, timeout=120), "add comment")
    if ok:
        _ok(_run("save", docx_path, timeout=60), "save")
        _ok(_run("close", docx_path, timeout=60), "close")
    return ok


def list_comments(docx_path: str) -> List[Dict[str, Any]]:
    r = _run("query", docx_path, "comment", "--json", timeout=120)
    if not r.get("ok"):
        return []
    try:
        data = json.loads(r.get("stdout") or "{}")
        return ((data.get("data") or {}).get("results")) or []
    except Exception:
        return []


# ---------------------------------------------------------------- 导出

def find_latest_step4_docx(project_id: str, out_dir: str = "") -> Optional[str]:
    d = step4_dir(project_id, out_dir=out_dir)
    if not d:
        return None
    side = _sidecar_path(d)
    if os.path.isfile(side):
        try:
            p = json.load(open(side, encoding="utf-8")).get("path") or ""
            if p and os.path.isfile(p):
                return p
        except Exception:
            pass
    try:
        cands = [os.path.join(d, f) for f in os.listdir(d)
                 if f.lower().endswith(".docx")]
        if not cands:
            return None
        cands.sort(key=lambda p: os.path.getmtime(p), reverse=True)
        return cands[0]
    except Exception:
        return None


def export_docx(src_path: str, dest_path: str) -> str:
    """释放 officecli 驻留并复制到目标路径（Step5 导出）。"""
    if not src_path or not os.path.isfile(src_path):
        raise FileNotFoundError(f"source docx not found: {src_path}")
    try:
        _run("close", src_path, timeout=60)
    except Exception:
        pass
    dest = str(dest_path)
    os.makedirs(os.path.dirname(dest) or ".", exist_ok=True)
    shutil.copy2(src_path, dest)
    return dest


# ================================================================ 正文同步
# 目标：交互（Rethink/局部重写/CGR 精修）的结果除批注外，**直接写回 Word 正文**。
# 手段：query paragraph → 定位段落（paraId 路径稳定）→ set（--find/--replace 优先，
# 保护锚定批注的段落）→ save/close。所有操作按文件加锁串行。

_LOCKS: Dict[str, threading.Lock] = {}
_LOCKS_GUARD = threading.Lock()


def _lock_for(path: str) -> threading.Lock:
    with _LOCKS_GUARD:
        lk = _LOCKS.get(path)
        if lk is None:
            lk = threading.Lock()
            _LOCKS[path] = lk
        return lk


_HEADING_ALIASES = {
    "abstract": ("abstract", "摘要"),
    "introduction": ("introduction", "引言", "背景"),
    "methods": ("method", "方法"),
    "results": ("result", "结果"),
    "discussion": ("discussion", "讨论"),
    "conclusion": ("conclusion", "结论"),
    "references": ("reference", "参考文献"),
}


def _norm_key(s: str) -> str:
    return re.sub(r"[\s#*:：\-、]+", "", str(s or "")).lower()


def _query_paragraphs(docx_path: str) -> List[Dict[str, Any]]:
    """正文段落（仅 /body/p[...]；含 styleId 与是否有图片）。"""
    r = _run("query", docx_path, "paragraph", "--json", timeout=120)
    if not r.get("ok"):
        return []
    try:
        data = json.loads(r.get("stdout") or "{}")
        results = ((data.get("data") or {}).get("results")) or []
    except Exception:
        return []
    out: List[Dict[str, Any]] = []
    for item in results:
        path = str(item.get("path") or "")
        if not path.startswith("/body/p["):
            continue
        fmt = item.get("format") if isinstance(item.get("format"), dict) else {}
        style_id = str(fmt.get("styleId") or item.get("style") or "")
        has_pic = any(isinstance(c, dict) and c.get("type") == "picture"
                      for c in (item.get("children") or []))
        out.append({
            "path": path,
            "text": str(item.get("text") or ""),
            "style_id": style_id,
            "is_heading": style_id.lower().startswith("heading"),
            "has_picture": has_pic,
        })
    return out


def _anchored_para_paths(docx_path: str) -> set:
    """批注锚定的段落路径集合（同步时保护，避免破坏批注范围）。"""
    anchored = set()
    for c in list_comments(docx_path):
        fmt = c.get("format") if isinstance(c.get("format"), dict) else {}
        a = str(fmt.get("anchoredTo") or "")
        if a:
            anchored.add(a)
    return anchored


def _heading_index(paras: List[Dict[str, Any]], heading: str) -> Optional[int]:
    target = _norm_key(heading)
    if not target:
        return None
    for i, p in enumerate(paras):
        if p["is_heading"] and target and target in _norm_key(p["text"]):
            return i
    for _key, aliases in _HEADING_ALIASES.items():
        if any(a in target for a in aliases):
            for i, p in enumerate(paras):
                if p["is_heading"] and any(a in _norm_key(p["text"]) for a in aliases):
                    return i
    return None


def _set_para_text(docx_path: str, para: Dict[str, Any], new_text: str) -> bool:
    """更新一个段落文本：优先 --find/--replace（保护锚定批注），兜底 --prop text=。"""
    one = str(new_text or "").replace("\n", "\x0b")
    old = str(para.get("text") or "")
    if old.strip():
        r = _run("set", docx_path, para["path"], "--find", old, "--replace", one, timeout=120)
        if r.get("ok"):
            return True
    return _ok(_run("set", docx_path, para["path"], "--prop", f"text={one}", timeout=120),
               "set text")


def sync_section_text(docx_path: str, heading: str, new_text: str) -> bool:
    """把某章节正文同步进 Word（heading 支持中英别名）。"""
    if not available() or not docx_path or not os.path.isfile(docx_path):
        return False
    content = str(new_text or "").strip()
    if not content:
        return False
    with _lock_for(docx_path):
        paras = _query_paragraphs(docx_path)
        h = _heading_index(paras, heading)
        if h is None:
            return False
        span: List[Dict[str, Any]] = []
        for p in paras[h + 1:]:
            if p["is_heading"]:
                break
            if p["text"].strip() and not p["has_picture"]:
                span.append(p)
        if not span:
            return False
        anchored = _anchored_para_paths(docx_path)
        blocks = _split_blocks(content)
        ok = True
        if len(blocks) == len(span):
            for p, b in zip(span, blocks):
                ok = _set_para_text(docx_path, p, b) and ok
        else:
            ok = _set_para_text(docx_path, span[0], content)
            for p in span[1:]:
                if p["path"] in anchored:
                    continue  # 不删除批注锚定段
                _ok(_run("remove", docx_path, p["path"], timeout=120), "remove para")
        if ok:
            _ok(_run("save", docx_path, timeout=60), "save")
            _ok(_run("close", docx_path, timeout=60), "close")
        return ok


def _sections_from_markdown(text: str) -> List[Any]:
    """极简 markdown 分节：返回 [(英文名, 内容)]（与构建端一致的关键词表）。"""
    t = str(text or "").replace("\r\n", "\n")
    sections: Dict[str, List[str]] = {}
    current: Optional[str] = None
    for line in t.splitlines():
        m = re.match(r"^#{1,4}\s+(.*?)\s*$", line)
        if m:
            low = m.group(1).lower()
            hit = None
            for key, aliases in _HEADING_ALIASES.items():
                if any(a in low for a in aliases):
                    hit = key
                    break
            if hit is not None:
                current = hit
                sections.setdefault(hit, [])
                continue
        if current is not None:
            sections.setdefault(current, []).append(line)
    out = []
    for key, _aliases in _HEADING_ALIASES.items():
        body = "\n".join(sections.get(key, [])).strip()
        if body:
            out.append((key, body))
    return out


def sync_full_text(docx_path: str, full_text: str) -> bool:
    """整篇更新：按 markdown 分节逐节同步（Rethink 场景）。"""
    parts = _sections_from_markdown(full_text)
    if not parts:
        return False
    any_ok = False
    for key, content in parts:
        if sync_section_text(docx_path, key, content):
            any_ok = True
    return any_ok


def sync_fragment_text(docx_path: str, old_fragment: str, new_fragment: str) -> bool:
    """片段级替换（CGR 精修）：定位含该片段的段落并原地替换。"""
    old = str(old_fragment or "").strip()
    new = str(new_fragment or "").strip()
    if not old or not new or not available() or not os.path.isfile(docx_path):
        return False
    with _lock_for(docx_path):
        paras = _query_paragraphs(docx_path)
        norm_old = _norm_key(old)[:80]
        target: Optional[Dict[str, Any]] = None
        for p in paras:
            if p["has_picture"]:
                continue
            if old in p["text"] or (norm_old and norm_old in _norm_key(p["text"])):
                target = p
                break
        if target is None:
            return False
        t = target["text"]
        ok = False
        if old in t:
            r = _run("set", docx_path, target["path"],
                     "--find", old, "--replace", new, timeout=120)
            if r.get("ok"):
                ok = True
        if not ok:
            r = _run("set", docx_path, target["path"], "--find", old,
                     "--replace", new, timeout=120)
            ok = bool(r.get("ok"))
        if ok:
            _ok(_run("save", docx_path, timeout=60), "save")
            _ok(_run("close", docx_path, timeout=60), "close")
        return ok


# ================================================================ 一稿一 docx（批量导出）

def _fallback_docx(out_path: str, title: str, sections: List[Any]) -> bool:
    """officecli 不可用时的兜底：python-docx 简版导出。"""
    try:
        from docx import Document  # type: ignore
        doc = Document()
        doc.add_heading(str(title or "Draft"), level=0)
        for name, content in sections:
            doc.add_heading(str(name).capitalize(), level=1)
            for para in _split_blocks(content):
                doc.add_paragraph(para)
        doc.save(out_path)
        return True
    except Exception as e:  # noqa: BLE001
        logger.warning("fallback docx failed: %s", e)
        return False


def build_docx_from_text(out_path: str, title: str, full_text: str,
                         image_paths: Optional[List[str]] = None) -> bool:
    """由一段 markdown 文本生成 docx（批量导出用；officecli 优先，python-docx 兜底）。"""
    sections = _sections_from_markdown(full_text)
    if not available():
        return _fallback_docx(out_path, title, sections)
    try:
        os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    except Exception:
        return False
    if not _ok(_run("create", out_path, timeout=120), "create"):
        return _fallback_docx(out_path, title, sections)
    _add_paragraph(out_path, str(title or "Draft"), style="Heading1", bold=True,
                   size="20", align="center")
    for key, content in sections:
        _add_paragraph(out_path, key.capitalize(), style="Heading1", bold=True, size="14")
        for b in _split_blocks(content):
            _add_paragraph(out_path, b, size="11")
    for i, img in enumerate((image_paths or [])[:8], 1):
        if os.path.isfile(str(img)):
            _add_picture(out_path, str(img), caption=f"图 {i}. {Path(img).name}")
    _ok(_run("validate", out_path, timeout=120), "validate")
    _ok(_run("save", out_path, timeout=120), "save")
    _ok(_run("close", out_path, timeout=120), "close")
    return os.path.isfile(out_path)


def export_drafts_batch(project_id: str, drafts: Optional[List[Dict[str, Any]]] = None,
                        dest_dir: str = "") -> List[str]:
    """Step5 批量导出：**一稿一 docx**。返回成功写出的文件路径列表。"""
    if drafts is None:
        try:
            from src.services.project_service import ProjectService
            drafts = ProjectService().get_project_drafts(str(project_id)) or []
        except Exception:
            drafts = []
    base = ""
    if dest_dir:
        base = str(dest_dir)
    else:
        storage = _storage_dir(str(project_id))
        if storage:
            base = os.path.join(storage, "exports", "step5")
    if not base:
        return []
    try:
        os.makedirs(base, exist_ok=True)
    except Exception:
        return []

    written: List[str] = []
    for d in drafts or []:
        if not isinstance(d, dict):
            continue
        content = str(d.get("content") or "").strip()
        if not content:
            continue
        title = str(d.get("title") or "Untitled Draft")
        did = str(d.get("id") or "")
        out = os.path.join(base, f"{_slug(title)}__{_slug(did) or 'draft'}.docx")
        try:
            if build_docx_from_text(out, title, content):
                written.append(out)
        except Exception as e:  # noqa: BLE001
            logger.warning("batch export failed for %s: %s", did, e)
    return written
